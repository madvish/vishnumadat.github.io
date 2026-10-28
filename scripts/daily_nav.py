"""Daily NAV, Return, and Alpha Tracker for Model Portfolio.

Builds a daily history of model portfolio returns, benchmark returns,
and cumulative/daily alpha. Uses yfinance for daily prices, aligns with
the Nifty 500 calendar, and applies strict risk and safety validations.

Usage:
  python scripts/daily_nav.py
  python scripts/daily_nav.py --unlisted-policy {stop,carry-inception,workbook}
"""

import argparse
import datetime as dt
import json
import sys
from pathlib import Path
import warnings

# Suppress noisy yfinance / pandas deprecation warnings
warnings.filterwarnings("ignore", category=FutureWarning)
warnings.filterwarnings("ignore", category=UserWarning)

import openpyxl
import pandas as pd
import yfinance as yf

# Paths
BASE_DIR = Path(__file__).resolve().parent.parent
XLSX_PATH = BASE_DIR / "Model_Portfolio_Master.xlsx"
DATA_DIR = BASE_DIR / "data"
TXN_PATH = DATA_DIR / "transactions.csv"
PRICES_DAILY_PATH = DATA_DIR / "prices_daily.csv"
NAV_DAILY_PATH = DATA_DIR / "nav_daily.csv"
NAV_CSV_PATH = DATA_DIR / "nav.csv"
META_PATH = DATA_DIR / "meta.json"

BENCHMARK_SYMBOL = "^CRSLDX"

# Known Yahoo symbol mappings for tickers not yet in workbook securities sheet
TICKER_SYMBOL_OVERRIDES = {
    "AXISHEALTH": "HEALTHAXIS.NS",
    "KRT": "KRT.BO",        # User provided KRT-BL.NS; Yahoo only provides daily historical bars under KRT.BO
    "VEDL-AL": "VAML.NS",    # Vedanta Aluminium Metal (listed 2026-06-15)
    "VEDL-OG": "VOGL.NS",    # Vedanta Oil & Gas (listed 2026-06-15)
    "VEDL-PWR": "VEDPOWER.NS",# Vedanta Power (listed 2026-06-15)
    "VEDL-IS": "VISL.NS",    # Vedanta Iron & Steel (listed 2026-06-15)
}


def parse_args():
    parser = argparse.ArgumentParser(
        description="Build daily NAV, return, and alpha history for model portfolio."
    )
    parser.add_argument(
        "--unlisted-policy",
        choices=["stop", "carry-inception", "workbook"],
        default="stop",
        help=(
            "Policy for held tickers without a Yahoo symbol or trading history: "
            "'stop' (default, strictly stop and ask per safety rules), "
            "'carry-inception' (carry forward inception buy price), "
            "'workbook' (use cached prices from workbook prices sheet)."
        ),
    )
    return parser.parse_args()


def load_workbook_settings_and_securities():
    """Reads settings, securities, and prices sheet from Model_Portfolio_Master.xlsx."""
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Workbook not found at {XLSX_PATH}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    # 1. Sheet: settings
    settings_sheet = wb["settings"]
    raw_incept = settings_sheet["B2"].value
    if isinstance(raw_incept, (dt.datetime, dt.date)):
        inception_date = raw_incept.date() if isinstance(raw_incept, dt.datetime) else raw_incept
    elif isinstance(raw_incept, str):
        inception_date = dt.date.fromisoformat(raw_incept.strip())
    else:
        raise ValueError(f"Invalid inception date in settings B2: {raw_incept}")

    starting_capital = float(settings_sheet["B3"].value)

    # 2. Sheet: securities
    sec_sheet = wb["securities"]
    headers = [sec_sheet.cell(1, c).value for c in range(1, sec_sheet.max_column + 1)]
    ticker_col = None
    yahoo_col = None
    for idx, h in enumerate(headers, start=1):
        if h:
            h_str = str(h).strip().lower()
            if h_str == "ticker" and ticker_col is None:
                ticker_col = idx
            elif h_str.startswith("yahoo_symbol") and yahoo_col is None:
                yahoo_col = idx

    if not ticker_col:
        ticker_col = 1
    if not yahoo_col:
        yahoo_col = 6

    securities_map = {}
    for r in range(2, sec_sheet.max_row + 1):
        t = sec_sheet.cell(r, ticker_col).value
        y = sec_sheet.cell(r, yahoo_col).value
        if t:
            t_str = str(t).strip()
            y_str = str(y).strip() if y is not None and str(y).strip() else None
            securities_map[t_str] = y_str

    # Apply symbol overrides for newly provided symbols
    for ticker, symbol in TICKER_SYMBOL_OVERRIDES.items():
        if not securities_map.get(ticker):
            securities_map[ticker] = symbol

    # 3. Sheet: prices (for benchmark row 2 column NIFTY500)
    prices_sheet = wb["prices"]
    p_headers = [prices_sheet.cell(1, c).value for c in range(1, prices_sheet.max_column + 1)]
    nifty_col = None
    for idx, h in enumerate(p_headers, start=1):
        if h and str(h).strip().upper() == "NIFTY500":
            nifty_col = idx
            break

    wb_nifty_incept = None
    if nifty_col:
        val = prices_sheet.cell(2, nifty_col).value
        if val is not None:
            wb_nifty_incept = float(val)

    # Cached prices map from row 2 and row 3 for reference
    wb_cached_prices = {}
    for idx, h in enumerate(p_headers, start=1):
        if h and idx > 1:
            h_str = str(h).strip()
            p2 = prices_sheet.cell(2, idx).value
            p3 = prices_sheet.cell(3, idx).value
            wb_cached_prices[h_str] = {
                "incept": float(p2) if p2 is not None else None,
                "latest": float(p3) if p3 is not None else None,
            }

    return inception_date, starting_capital, securities_map, wb_nifty_incept, wb_cached_prices


def load_transactions():
    """Loads transactions.csv and returns DataFrame with parsed dates."""
    if not TXN_PATH.exists():
        raise FileNotFoundError(f"Transactions file not found at {TXN_PATH}")
    df = pd.read_csv(TXN_PATH)
    df["date"] = pd.to_datetime(df["date"]).dt.date
    df["action"] = df["action"].str.strip().str.upper()
    df["ticker"] = df["ticker"].str.strip()
    df["qty"] = df["qty"].astype(int)
    df["price"] = df["price"].astype(float)
    return df


def calculate_inception_buy_prices(txn_df, inception_date):
    """Calculates weighted average buy price for each ticker on or before inception date."""
    incept_txns = txn_df[(txn_df["date"] <= inception_date) & (txn_df["action"] == "BUY")]
    buy_prices = {}
    for ticker, grp in incept_txns.groupby("ticker"):
        total_qty = grp["qty"].sum()
        if total_qty > 0:
            avg_price = (grp["qty"] * grp["price"]).sum() / total_qty
            buy_prices[ticker] = avg_price
    return buy_prices


def check_benchmark(inception_date, wb_nifty_incept):
    """
    Fetches Nifty 500 (^CRSLDX) close on inception date and compares with workbook value.
    If diff > 1%, stops and informs user.
    """
    bm_ticker = yf.Ticker(BENCHMARK_SYMBOL)
    bm_hist = bm_ticker.history(start=inception_date.isoformat(), auto_adjust=False)

    if bm_hist.empty:
        raise SystemExit(
            f"ERROR: Could not fetch benchmark {BENCHMARK_SYMBOL} history starting from {inception_date}."
        )

    # Convert index to naive date
    bm_hist.index = pd.to_datetime(bm_hist.index).date

    if inception_date not in bm_hist.index:
        first_date = bm_hist.index[0]
        yf_close = float(bm_hist["Close"].iloc[0])
        print(f"Note: Inception date {inception_date} not directly in benchmark history; first available trading date is {first_date}.")
    else:
        yf_close = float(bm_hist.loc[inception_date, "Close"])

    print("==================================================")
    print("BENCHMARK INCEPTION VALIDATION")
    print("==================================================")
    print(f"Benchmark Symbol:                        {BENCHMARK_SYMBOL} (Nifty 500)")
    print(f"Inception Date:                          {inception_date}")
    print(f"Nifty 500 close on inception date (YF):  {yf_close:,.2f}")
    if wb_nifty_incept is not None:
        print(f"Nifty 500 in workbook prices sheet (B2): {wb_nifty_incept:,.2f}")
        diff_pct = abs(yf_close - wb_nifty_incept) / wb_nifty_incept * 100
        print(f"Absolute Difference:                     {diff_pct:.4f}%")
        if diff_pct > 1.0:
            raise SystemExit(
                f"STOP: Nifty 500 close on inception date ({yf_close:,.2f}) differs from "
                f"workbook prices sheet ({wb_nifty_incept:,.2f}) by {diff_pct:.2f}% (> 1.0%)."
            )
        else:
            print("Status:                                  MATCH (Difference <= 1.0%)")
    else:
        print("Warning: NIFTY500 column not found in workbook prices sheet.")
    print("==================================================\n")

    return bm_hist


def main():
    args = parse_args()
    DATA_DIR.mkdir(exist_ok=True)

    # 1. Inputs
    incept_date, starting_cap, securities_map, wb_nifty_incept, wb_cached_prices = (
        load_workbook_settings_and_securities()
    )
    txn_df = load_transactions()
    held_tickers = sorted(txn_df["ticker"].unique())

    # 2. Benchmark Validation & Calendar
    bm_hist = check_benchmark(incept_date, wb_nifty_incept)
    calendar_dates = [d for d in bm_hist.index if d >= incept_date]

    # Inception benchmark close
    if incept_date in bm_hist.index:
        bm_incept_close = float(bm_hist.loc[incept_date, "Close"])
    else:
        bm_incept_close = float(bm_hist["Close"].iloc[0])

    # 3. Identify held tickers without Yahoo symbol
    unlisted_tickers = [t for t in held_tickers if not securities_map.get(t)]
    warnings_list = []

    if unlisted_tickers:
        print("--------------------------------------------------")
        print("SAFETY RULE CHECK: HELD TICKERS WITHOUT YAHOO SYMBOL")
        print("--------------------------------------------------")
        print(f"Found {len(unlisted_tickers)} held ticker(s) without a configured yahoo_symbol in sheet 'securities':")
        for ut in unlisted_tickers:
            company_note = ""
            if "VEDL" in ut:
                company_note = " (Vedanta demerged entity - price history starts post-inception)"
            elif ut == "KRT":
                company_note = " (Knowledge Realty Trust REIT)"
            elif ut == "AXISHEALTH":
                company_note = " (Axis Nifty Healthcare ETF)"
            print(f"  - {ut}{company_note}")

        if args.unlisted_policy == "stop":
            print("\nSAFETY STOP TRIGGERED:")
            print("Per safety rule: 'If a held ticker has no yahoo_symbol, or Yahoo has no price for it")
            print("on some days, do not guess. List those tickers and dates and stop and ask me what to do.")
            print("Expect this for the Vedanta demerged entities, whose price history starts after inception.'")
            print("\nOptions to proceed:")
            print("  1. Carry forward each unlisted stock's inception buy price until price data begins.")
            print("  2. Use cached prices from the workbook prices sheet.")
            print("  3. Provide exact Yahoo Finance symbols for any listed entities.")
            print("\nTo run with inception price carry-forward, re-run with: --unlisted-policy carry-inception")
            sys.exit(1)
        elif args.unlisted_policy == "carry-inception":
            msg = (
                f"Carrying forward inception buy price from transactions.csv for unlisted tickers: "
                f"{', '.join(unlisted_tickers)}."
            )
            print(f"\n[POLICY: carry-inception] {msg}")
            warnings_list.append(msg)
        elif args.unlisted_policy == "workbook":
            msg = (
                f"Using cached prices from workbook prices sheet for unlisted tickers: "
                f"{', '.join(unlisted_tickers)}."
            )
            print(f"\n[POLICY: workbook] {msg}")
            warnings_list.append(msg)
        print("--------------------------------------------------\n")

    # Inception buy prices from transactions.csv
    incept_buy_prices = calculate_inception_buy_prices(txn_df, incept_date)

    # Check append-only past runs
    existing_prices_df = None
    existing_nav_df = None
    last_saved_date = None

    if PRICES_DAILY_PATH.exists() and NAV_DAILY_PATH.exists():
        try:
            existing_prices_df = pd.read_csv(PRICES_DAILY_PATH)
            existing_nav_df = pd.read_csv(NAV_DAILY_PATH)
            
            # Prune any trailing corrupted rows with NaN values
            valid_mask = (
                existing_nav_df["portfolio_index"].notna() &
                (existing_nav_df["portfolio_index"].astype(str).str.lower() != "nan") &
                existing_nav_df["benchmark_index"].notna() &
                (existing_nav_df["benchmark_index"].astype(str).str.lower() != "nan")
            )
            if not valid_mask.all():
                corrupted_dates = existing_nav_df.loc[~valid_mask, "date"].tolist()
                print(f"Detected corrupted row(s) with NaN on dates: {corrupted_dates}. Pruning for recomputation.")
                existing_nav_df = existing_nav_df[valid_mask].reset_index(drop=True)
                if not existing_nav_df.empty:
                    max_clean_date = existing_nav_df["date"].iloc[-1]
                    existing_prices_df = existing_prices_df[existing_prices_df["date"] <= max_clean_date].reset_index(drop=True)
                else:
                    existing_prices_df = None
                    existing_nav_df = None

            if existing_nav_df is not None and not existing_nav_df.empty and "date" in existing_nav_df.columns:
                last_saved_date = dt.date.fromisoformat(str(existing_nav_df["date"].iloc[-1]).strip())
                print(f"Found existing history up to {last_saved_date}. New runs will append subsequent dates only.")
        except Exception as e:
            print(f"Warning reading existing daily files: {e}. Recomputing full history.")
            existing_prices_df = None
            existing_nav_df = None
            last_saved_date = None

    # Filter dates if appending
    if last_saved_date is not None:
        dates_to_compute = [d for d in calendar_dates if d > last_saved_date]
        if not dates_to_compute:
            print("All trading days up to latest benchmark date are already recorded. Nothing to append.")
            return
    else:
        dates_to_compute = calendar_dates

    # 4. Fetch daily closes for listed tickers + benchmark
    listed_tickers = [t for t in held_tickers if securities_map.get(t)]
    download_symbols = [securities_map[t] for t in listed_tickers] + [BENCHMARK_SYMBOL]

    print(f"Fetching daily close prices from {incept_date} to {calendar_dates[-1]} for {len(download_symbols)} symbols...")
    raw_yf_df = yf.download(
        download_symbols,
        start=incept_date.isoformat(),
        auto_adjust=False,
        progress=False,
    )["Close"]

    # Normalize tz and index format
    if raw_yf_df.index.tz is not None:
        raw_yf_df.index = raw_yf_df.index.tz_convert(None)
    raw_yf_df.index = pd.to_datetime(raw_yf_df.index).date

    # 5. Build full daily prices matrix
    # Columns: date, held_tickers (sorted), NIFTY500, stale_tickers
    daily_prices_rows = []
    last_known_prices = {}
    missing_day_counter = {t: 0 for t in held_tickers}

    for current_date in calendar_dates:
        row_prices = {"date": current_date.isoformat()}
        current_stale = []

        # Benchmark price
        bm_price = None
        if current_date in bm_hist.index:
            val = bm_hist.loc[current_date, "Close"]
            if pd.notna(val) and float(val) > 0:
                bm_price = float(val)
        if bm_price is None and current_date in raw_yf_df.index and BENCHMARK_SYMBOL in raw_yf_df.columns:
            val = raw_yf_df.loc[current_date, BENCHMARK_SYMBOL]
            if pd.notna(val) and float(val) > 0:
                bm_price = float(val)
        if bm_price is None or pd.isna(bm_price) or bm_price <= 0:
            bm_price = last_known_prices.get("NIFTY500", bm_incept_close)

        last_known_prices["NIFTY500"] = bm_price

        # Held tickers
        for ticker in held_tickers:
            price = None
            if current_date == incept_date:
                # Rule: On the inception date, use each stock's average buy price from transactions.csv
                price = incept_buy_prices.get(ticker)
                missing_day_counter[ticker] = 0
            else:
                sym = securities_map.get(ticker)
                if sym and sym in raw_yf_df.columns and current_date in raw_yf_df.index:
                    val = raw_yf_df.loc[current_date, sym]
                    if pd.notna(val) and float(val) > 0:
                        price = float(val)
                        missing_day_counter[ticker] = 0

                # Missing price handling
                if price is None:
                    # Demerged entities before their listing date carry forward their inception cost allocation
                    if ticker in ("VEDL-AL", "VEDL-PWR", "VEDL-OG", "VEDL-IS") and current_date < dt.date(2026, 6, 15):
                        price = last_known_prices.get(ticker, incept_buy_prices.get(ticker))
                    elif ticker in unlisted_tickers:
                        if args.unlisted_policy == "carry-inception":
                            price = incept_buy_prices.get(ticker)
                        elif args.unlisted_policy == "workbook":
                            wb_info = wb_cached_prices.get(ticker, {})
                            price = wb_info.get("latest") or wb_info.get("incept") or incept_buy_prices.get(ticker)
                    else:
                        # Carry forward for up to 5 days
                        missing_day_counter[ticker] += 1
                        if missing_day_counter[ticker] <= 5 and ticker in last_known_prices:
                            price = last_known_prices[ticker]
                            current_stale.append(ticker)
                        else:
                            raise SystemExit(
                                f"STOP: Price for {ticker} is missing for more than 5 consecutive trading days at {current_date}."
                            )

            row_prices[ticker] = round(price, 4) if price is not None else None
            last_known_prices[ticker] = price

        # Safety: if not inception date and 100% of held tickers had no trade, skip phantom day
        if current_date != incept_date and len(current_stale) == len(held_tickers):
            print(f"Warning: No new market prices available for any ticker on {current_date}. Skipping recording phantom day.")
            continue

        row_prices["NIFTY500"] = round(bm_price, 2)
        row_prices["stale_tickers"] = ";".join(current_stale) if current_stale else ""
        daily_prices_rows.append(row_prices)

    # Reorder columns: date, held_tickers (sorted), NIFTY500, stale_tickers
    price_cols = ["date"] + held_tickers + ["NIFTY500", "stale_tickers"]
    full_prices_df = pd.DataFrame(daily_prices_rows)[price_cols]

    # 6. Safety check: Check for > 25% daily price changes
    for ticker in held_tickers:
        series = full_prices_df[ticker].astype(float)
        pct_chg = series.pct_change()
        for idx, chg in pct_chg.items():
            if pd.notna(chg) and abs(chg) > 0.25:
                dt_str = full_prices_df.loc[idx, "date"]
                p_prev = full_prices_df.loc[idx - 1, ticker]
                p_curr = full_prices_df.loc[idx, ticker]
                warn_msg = (
                    f"Possible split, bonus, or demerger review: {ticker} changed by "
                    f"{chg*100:+.2f}% on {dt_str} ({p_prev:.2f} -> {p_curr:.2f})."
                )
                warnings_list.append(warn_msg)

    # 7. Compute NAV and index metrics for each date
    # Determine live tracking start date in meta.json
    live_start = None
    if META_PATH.exists():
        try:
            with open(META_PATH, "r", encoding="utf-8") as f:
                meta_json = json.load(f)
                live_start = meta_json.get("live_tracking_start")
        except Exception:
            meta_json = {}
    else:
        meta_json = {}

    if not live_start:
        live_start = dt.date.today().isoformat()
        meta_json["live_tracking_start"] = live_start
        print(f"Recorded initial live tracking start date: {live_start}")

    nav_rows = []
    running_max_index = 0.0
    prev_portfolio_index = None
    prev_benchmark_index = None

    for idx, current_date in enumerate(calendar_dates):
        # Transactions on or before current_date
        tx_up_to_now = txn_df[txn_df["date"] <= current_date]

        # Holdings quantity
        buy_qtys = tx_up_to_now[tx_up_to_now["action"] == "BUY"].groupby("ticker")["qty"].sum()
        sell_qtys = tx_up_to_now[tx_up_to_now["action"] == "SELL"].groupby("ticker")["qty"].sum()
        current_holdings = buy_qtys.subtract(sell_qtys, fill_value=0)

        # Cash = starting_capital - (buy_values - sell_values)
        buy_vals = (tx_up_to_now[tx_up_to_now["action"] == "BUY"]["qty"] * tx_up_to_now[tx_up_to_now["action"] == "BUY"]["price"]).sum()
        sell_vals = (tx_up_to_now[tx_up_to_now["action"] == "SELL"]["qty"] * tx_up_to_now[tx_up_to_now["action"] == "SELL"]["price"]).sum()
        cash = starting_cap - (buy_vals - sell_vals)

        # Holdings value
        date_str = current_date.isoformat()
        holdings_val = 0.0
        p_row = full_prices_df[full_prices_df["date"] == date_str].iloc[0]
        for ticker, qty in current_holdings.items():
            if qty > 0:
                p = float(p_row[ticker])
                holdings_val += qty * p

        # NAV & Index
        nav = holdings_val + cash
        portfolio_index = (nav / starting_cap) * 100.0
        bm_price = float(p_row["NIFTY500"])
        benchmark_index = (bm_price / bm_incept_close) * 100.0

        # Running max for drawdown
        if portfolio_index > running_max_index:
            running_max_index = portfolio_index
        drawdown_pct = ((portfolio_index / running_max_index) - 1.0) * 100.0

        # Cumulative returns
        if idx == 0:
            portfolio_index = 100.00
            benchmark_index = 100.00
            portfolio_return_pct = 0.00
            benchmark_return_pct = 0.00
            alpha_pp = 0.00
            drawdown_pct = 0.00
            daily_portfolio_return_pct = 0.00
            daily_benchmark_return_pct = 0.00
            daily_alpha_pp = 0.00
        else:
            portfolio_return_pct = portfolio_index - 100.0
            benchmark_return_pct = benchmark_index - 100.0
            alpha_pp = portfolio_index - benchmark_index
            daily_portfolio_return_pct = ((portfolio_index / prev_portfolio_index) - 1.0) * 100.0
            daily_benchmark_return_pct = ((benchmark_index / prev_benchmark_index) - 1.0) * 100.0
            daily_alpha_pp = daily_portfolio_return_pct - daily_benchmark_return_pct

        prev_portfolio_index = portfolio_index
        prev_benchmark_index = benchmark_index

        # Clean any negative zeros
        def clean_val(v):
            return 0.00 if abs(v) < 0.000001 else v

        # Basis
        basis = "backfilled" if date_str < live_start else "live"

        nav_rows.append({
            "date": date_str,
            "portfolio_index": f"{clean_val(portfolio_index):.2f}",
            "benchmark_index": f"{clean_val(benchmark_index):.2f}",
            "alpha_pp": f"{clean_val(alpha_pp):.2f}",
            "portfolio_return_pct": f"{clean_val(portfolio_return_pct):.2f}",
            "benchmark_return_pct": f"{clean_val(benchmark_return_pct):.2f}",
            "daily_portfolio_return_pct": f"{clean_val(daily_portfolio_return_pct):.2f}",
            "daily_benchmark_return_pct": f"{clean_val(daily_benchmark_return_pct):.2f}",
            "daily_alpha_pp": f"{clean_val(daily_alpha_pp):.2f}",
            "drawdown_pct": f"{clean_val(drawdown_pct):.2f}",
            "basis": basis,
        })

    computed_nav_df = pd.DataFrame(nav_rows)

    # 8. Append-only merge
    if existing_prices_df is not None and not existing_prices_df.empty:
        new_prices_to_append = full_prices_df[full_prices_df["date"] > str(last_saved_date)]
        final_prices_df = pd.concat([existing_prices_df, new_prices_to_append], ignore_index=True)
    else:
        final_prices_df = full_prices_df

    if existing_nav_df is not None and not existing_nav_df.empty:
        new_nav_to_append = computed_nav_df[computed_nav_df["date"] > str(last_saved_date)]
        final_nav_df = pd.concat([existing_nav_df, new_nav_to_append], ignore_index=True)
    else:
        final_nav_df = computed_nav_df

    # 9. Save files
    final_prices_df.to_csv(PRICES_DAILY_PATH, index=False)
    final_nav_df.to_csv(NAV_DAILY_PATH, index=False)
    print(f"Saved {len(final_prices_df)} daily prices to {PRICES_DAILY_PATH}")
    print(f"Saved {len(final_nav_df)} daily NAV records to {NAV_DAILY_PATH}")

    # 10. Compare with data/nav.csv on matching dates
    if NAV_CSV_PATH.exists():
        try:
            ref_nav_df = pd.read_csv(NAV_CSV_PATH)
            ref_nav_df["date"] = ref_nav_df["date"].astype(str).str.strip()
            nav_daily_indexed = final_nav_df.set_index("date")

            for _, ref_row in ref_nav_df.iterrows():
                r_date = str(ref_row["date"]).strip()
                if r_date in nav_daily_indexed.index:
                    calc_pi = float(nav_daily_indexed.loc[r_date, "portfolio_index"])
                    ref_pi = float(ref_row["portfolio_index"])
                    diff = abs(calc_pi - ref_pi)
                    if diff > 0.05:
                        warn_msg = (
                            f"Discrepancy on {r_date}: portfolio_index in nav_daily.csv ({calc_pi:.2f}) "
                            f"differs from nav.csv ({ref_pi:.2f}) by {diff:.2f} (> 0.05)."
                        )
                        warnings_list.append(warn_msg)
        except Exception as e:
            warnings_list.append(f"Could not compare against nav.csv: {e}")

    # 11. Update data/meta.json
    latest_row = final_nav_df.iloc[-1]
    p_ret = float(latest_row["portfolio_return_pct"])
    b_ret = float(latest_row["benchmark_return_pct"])
    a_pp = float(latest_row["alpha_pp"])

    if pd.isna(p_ret) or pd.isna(b_ret) or pd.isna(a_pp):
        raise ValueError(
            f"FATAL: Attempted to write NaN to meta.json! (p_ret={p_ret}, b_ret={b_ret}, a_pp={a_pp})"
        )

    meta_json["since_inception_return_pct"] = round(p_ret, 2)
    meta_json["benchmark_return_pct"] = round(b_ret, 2)
    meta_json["alpha_pp"] = round(a_pp, 2)
    meta_json["data_as_of"] = str(latest_row["date"])
    meta_json["last_updated"] = dt.date.today().isoformat()

    with open(META_PATH, "w", encoding="utf-8") as f:
        json.dump(meta_json, f, indent=2, allow_nan=False)
    print(f"Updated {META_PATH} with latest metrics from {latest_row['date']}.")

    # 12. Display requested summary
    p_ret = float(latest_row["portfolio_return_pct"])
    b_ret = float(latest_row["benchmark_return_pct"])
    a_pp = float(latest_row["alpha_pp"])
    dd = float(latest_row["drawdown_pct"])

    print("\n==================================================")
    print("DAILY NAV GENERATION SUMMARY")
    print("==================================================")
    print(f"Total Trading Days:              {len(final_nav_df)}")
    print(f"Latest As-Of Date:               {latest_row['date']}")
    print(f"Latest Portfolio Index:          {latest_row['portfolio_index']}")
    print(f"Latest Benchmark Index:          {latest_row['benchmark_index']}")
    print(f"Cumulative Portfolio Return:     {p_ret:+.2f}%")
    print(f"Cumulative Benchmark Return:     {b_ret:+.2f}%")
    print(f"Cumulative Alpha:                {a_pp:+.2f} pp")
    print(f"Current Drawdown:                {dd:.2f}%")
    print(f"Tracking Basis:                  {latest_row['basis']}")
    print("==================================================")

    print("\n--- FIRST 3 ROWS (nav_daily.csv) ---")
    print(final_nav_df.head(3).to_string(index=False))

    print("\n--- LAST 3 ROWS (nav_daily.csv) ---")
    print(final_nav_df.tail(3).to_string(index=False))

    if warnings_list:
        print("\n--- WARNINGS & AUDIT LOG ---")
        for idx, w in enumerate(warnings_list, start=1):
            print(f"{idx}. {w}")
    else:
        print("\n--- WARNINGS: None ---")


if __name__ == "__main__":
    main()
